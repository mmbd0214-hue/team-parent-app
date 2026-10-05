"""Integration tests only target an explicitly supplied, disposable localhost DB."""
import os
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
from io import BytesIO
from urllib.parse import urlparse
import time
from types import SimpleNamespace
import jwt

import psycopg
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

URL = os.getenv('TEST_DATABASE_URL', '')
if URL:
    target = urlparse(URL)
    if target.hostname != '127.0.0.1' or target.port != 55432 or not target.path.endswith('_test'):
        raise RuntimeError('Tests refuse non-local or non-disposable databases')
os.environ['DATABASE_URL'] = URL or 'postgresql://invalid/unused'
os.environ['APP_ENV'] = 'development'
os.environ['MOCK_LOGIN'] = '1'
os.environ['ADMIN_PASSWORD'] = 'test-only-strong-password'
os.environ['APP_BASE_URL'] = 'https://team.test'

import app as backend
from session_security import issue_session, token_hash
from mobile_api import taipei_today, validate_report


@pytest.fixture
def seeded():
    if not URL:
        pytest.skip('Set TEST_DATABASE_URL to an isolated local database')
    with backend.db() as conn, conn.cursor() as cur:
        cur.execute('TRUNCATE parents,players,events,announcements,api_rate_limits,apple_challenges,provider_revocations RESTART IDENTITY CASCADE')
        for name in ['A', 'B', 'outsider']:
            cur.execute('INSERT INTO parents(line_user_id,display_name,is_admin) VALUES(%s,%s,%s)', ('line-'+name, name, name == 'A'))
        for name in ['childA', 'childB']:
            cur.execute("INSERT INTO players(name,team,number,bind_code) VALUES(%s,'U10','01',%s)", (name, 'AAAAAA' if name == 'childA' else 'BBBBBB'))
        cur.execute('INSERT INTO parent_players VALUES(1,1),(1,2),(2,1)')
        cur.execute("INSERT INTO events(title,event_date,location,response_deadline) VALUES('Practice',%s,'park',%s)", (taipei_today(), taipei_today()))
        cur.execute('INSERT INTO event_players VALUES(1,1)')
        cur.execute("INSERT INTO payments(player_id,title,amount) VALUES(1,'fee',500)")
    tokens = [issue_session(backend.db, i)['access_token'] for i in [1,2,3]]
    with TestClient(backend.app, base_url='https://team.test') as client:
        yield client, tokens


def headers(token):
    return {'Authorization': 'Bearer '+token}


def reply(version=0):
    return {'player_id': 1, 'attendance_status': 'attend', 'practice_duration': 'morning_leave',
            'attendance_note': 'afternoon only', 'leave_reason': '', 'expected_version': version}


def report(version=1):
    return {'payment_method': 'transfer', 'transfer_date': '2026-10-02', 'account_last5': '00123', 'expected_version': version}


@pytest.mark.parametrize('value', ['parent:1', 'parent-admin:1', 'admin:test-only-strong-password'])
def test_legacy_credentials_rejected(seeded, value):
    client, _ = seeded
    assert client.get('/api/me', headers=headers(value)).status_code == 401
    assert client.get('/api/admin/payments', headers=headers(value)).status_code == 401


def test_refresh_rotation_and_replay_revokes_session(seeded):
    client, _ = seeded
    initial = issue_session(backend.db, 1)
    changed = client.post('/api/auth/refresh', json={'refresh_token': initial['refresh_token']})
    assert changed.status_code == 200
    token = changed.json()['access_token']
    assert token != initial['access_token']
    assert client.get('/api/me', headers=headers(initial['access_token'])).status_code == 401
    assert client.get('/api/me', headers=headers(token)).status_code == 200
    assert client.post('/api/auth/refresh', json={'refresh_token': initial['refresh_token']}).status_code == 401
    assert client.get('/api/me', headers=headers(token)).status_code == 401


def test_logout_revokes_and_is_idempotent(seeded):
    client, tokens = seeded
    assert client.post('/api/auth/logout', headers=headers(tokens[0])).status_code == 200
    assert client.post('/api/auth/logout', headers=headers(tokens[0])).status_code == 200
    assert client.get('/api/me', headers=headers(tokens[0])).status_code == 401


def test_player_scope_and_authorization(seeded):
    client, tokens = seeded
    assert len(client.get('/api/events?player_id=1', headers=headers(tokens[0])).json()) == 1
    assert client.get('/api/events?player_id=2', headers=headers(tokens[0])).json() == []
    assert client.get('/api/players/1/payments', headers=headers(tokens[2])).status_code == 403
    assert client.get('/api/events/1?player_id=2', headers=headers(tokens[0])).status_code == 403
    body = reply(); body['player_id'] = 2
    assert client.put('/api/events/1/attendance', json=body, headers=headers(tokens[0])).status_code == 403
    assert client.get('/api/events/1/attendance-summary', headers=headers(tokens[2])).status_code == 403


def test_attendance_conflict_and_old_meals_preserved(seeded):
    client, tokens = seeded
    with backend.db() as conn, conn.cursor() as cur:
        cur.execute("INSERT INTO attendance(event_id,player_id,attendance_status,player_meals,parent_meals) VALUES(1,1,'attend',2,3)")
    first = client.put('/api/events/1/attendance', json=reply(1), headers=headers(tokens[0]))
    assert first.status_code == 200
    saved = first.json()['attendance']
    assert saved['player_meals'] == 2 and saved['parent_meals'] == 3
    assert saved['version'] == 2 and saved['attendance_note'] == 'afternoon only'
    assert client.put('/api/events/1/attendance', json=reply(1), headers=headers(tokens[1])).status_code == 409


def test_concurrent_first_attendance_is_not_silently_overwritten(seeded):
    client, tokens = seeded
    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(lambda token: client.put('/api/events/1/attendance', json=reply(), headers=headers(token)), tokens[:2]))
    assert sorted(r.status_code for r in responses) == [200, 409]


@pytest.mark.parametrize('column,value', [('status', 'closed'), ('survey_enabled', False), ('response_deadline', '2020-01-01')])
def test_closed_notice_or_expired_rejects_reply(seeded, column, value):
    client, tokens = seeded
    with backend.db() as conn, conn.cursor() as cur:
        cur.execute(psycopg.sql.SQL('UPDATE events SET {}=%s WHERE id=1').format(psycopg.sql.Identifier(column)), (value,))
    assert client.put('/api/events/1/attendance', json=reply(), headers=headers(tokens[0])).status_code == 403
    detail = client.get('/api/events/1?player_id=1', headers=headers(tokens[0])).json()
    assert detail['eligibility']['can_reply'] is False


def test_paid_cannot_be_changed_back_to_pending(seeded):
    client, tokens = seeded
    assert client.put('/api/payments/1/transfer', json=report(), headers=headers(tokens[0])).status_code == 200
    paid = client.put('/api/admin/payments/1/status?status=paid&expected_version=2', headers=headers(tokens[0]))
    assert paid.status_code == 200
    assert client.put('/api/payments/1/transfer', json=report(2), headers=headers(tokens[0])).status_code == 409
    assert client.get('/api/players/1/payments', headers=headers(tokens[0])).json()[0]['status'] == 'paid'


def test_report_validation_and_leading_zeros(seeded):
    client, tokens = seeded
    invalid = report(); invalid['account_last5'] = '１２３４５'
    assert client.put('/api/payments/1/transfer', json=invalid, headers=headers(tokens[0])).status_code == 422
    invalid = report(); invalid['transfer_date'] = 'not-a-date'
    assert client.put('/api/payments/1/transfer', json=invalid, headers=headers(tokens[0])).status_code == 422
    response = client.put('/api/payments/1/transfer', json=report(), headers=headers(tokens[0]))
    assert response.json()['transfer_account_last5'] == '00123'


def test_admin_role_revocation_immediate(seeded):
    client, tokens = seeded
    assert client.get('/api/admin/payments', headers=headers(tokens[0])).status_code == 200
    with backend.db() as conn, conn.cursor() as cur:
        cur.execute('UPDATE parents SET is_admin=FALSE WHERE id=1')
    assert client.get('/api/admin/payments', headers=headers(tokens[0])).status_code == 403
    assert client.post('/api/auth/web-handoff', headers=headers(tokens[0])).status_code == 403


def test_handoff_one_use_and_cookie_csrf(seeded):
    client, tokens = seeded
    url = client.post('/api/auth/web-handoff', headers=headers(tokens[0])).json()['url']
    assert client.get(url, follow_redirects=False).status_code == 303
    assert client.get(url, follow_redirects=False).status_code == 401
    session = client.get('/api/admin/session').json()
    path = '/api/admin/payments/1/status?status=paid&expected_version=1'
    assert client.put(path).status_code == 403
    assert client.put(path, headers={'X-CSRF-Token': session['csrf_token']}).status_code == 200


def test_content_seen_never_regresses(seeded):
    client, tokens = seeded
    for value in [12, 5]:
        response = client.post('/api/content-seen', headers=headers(tokens[0]), json={'kind': 'payment', 'scope': '1', 'last_seen_id': value})
        assert response.status_code == 200
        assert response.json()['last_seen_id'] == 12


def test_excel_is_text_and_null_filters_match(seeded):
    client, tokens = seeded
    with backend.db() as conn, conn.cursor() as cur:
        cur.execute("UPDATE payments SET title='=1+1',note='',transfer_account_last5='00123',payment_method='transfer' WHERE id=1")
        cur.execute("INSERT INTO payments(player_id,title,amount,due_date) VALUES(2,'=1+1',500,'2027-01-01')")
    response = client.get('/api/admin/payments/export.xlsx?title=%3D1%2B1&amount=500&due_date=&note=', headers=headers(tokens[0]))
    assert response.status_code == 200
    ws = load_workbook(BytesIO(response.content)).active
    assert ws.max_row == 2
    assert ws['A2'].value == '=1+1' and ws['A2'].data_type == 's'
    assert ws['I2'].value == '00123'


def test_account_deletion_preserves_shared_player(seeded):
    client, tokens = seeded
    result = client.post('/api/me/deletion-request', headers=headers(tokens[0]), json={'confirmation': '刪除帳號'})
    assert result.status_code == 200
    assert client.get('/api/me', headers=headers(tokens[0])).status_code == 401
    assert client.get('/api/me', headers=headers(tokens[1])).json()['players'][0]['id'] == 1
    assert client.get('/api/players/1/payments', headers=headers(tokens[1])).status_code == 200


def test_migration_idempotent_and_restart_does_not_assign_empty_events(seeded):
    from pathlib import Path
    client, tokens = seeded
    with backend.db() as conn, conn.cursor() as cur:
        cur.execute("INSERT INTO events(title,event_date,location) VALUES('Empty',%s,'park')", (taipei_today(),))
        cur.execute(Path('migrations/001_mobile.sql').read_text(encoding='utf-8'))
    backend.startup()
    with backend.db() as conn, conn.cursor() as cur:
        cur.execute('SELECT COUNT(*) n FROM event_players WHERE event_id=2')
        assert cur.fetchone()['n'] == 0


def test_date_and_cash_validation():
    assert validate_report('cash', '2026-10-02', '00123')[1] == ''
    assert token_hash('private-secret') != 'private-secret'


def test_announcements_filter_before_pagination(seeded):
    client, tokens = seeded
    with backend.db() as conn, conn.cursor() as cur:
        cur.execute("INSERT INTO announcements(title,message_text) VALUES('visible','news')")
        for _ in range(305):
            cur.execute("INSERT INTO announcements(message_text,target_type,target_values) VALUES('not mine','team','[\"U15\"]')")
    page = client.get('/api/announcements?limit=100', headers=headers(tokens[0])).json()
    assert [r['title'] for r in page['items']] == ['visible']
    assert page['next_cursor'] is None


def test_inactive_player_rejected(seeded):
    client, tokens = seeded
    with backend.db() as conn, conn.cursor() as cur:
        cur.execute('UPDATE players SET active=FALSE WHERE id=1')
    for path in ['/api/events?player_id=1', '/api/players/1/attendance', '/api/players/1/payments']:
        assert client.get(path, headers=headers(tokens[0])).status_code == 403


def test_deleting_paid_payment_requires_reason_and_audit(seeded):
    client, tokens = seeded
    with backend.db() as conn, conn.cursor() as cur:
        cur.execute("UPDATE payments SET status='paid' WHERE id=1")
    assert client.delete('/api/admin/payments/1?expected_version=1', headers=headers(tokens[0])).status_code == 422
    assert client.delete('/api/admin/payments/1?expected_version=1&reason=duplicate', headers=headers(tokens[0])).status_code == 200
    with backend.db() as conn, conn.cursor() as cur:
        cur.execute("SELECT details FROM audit_logs WHERE action='payment_deleted'")
        assert cur.fetchone()['details']['reason'] == 'duplicate'


def test_apple_nonce_and_signature_then_encrypted_revocation(seeded, monkeypatch):
    from cryptography.hazmat.primitives.asymmetric import rsa
    from cryptography.fernet import Fernet
    import mobile_api
    import apple_identity
    client, _ = seeded
    monkeypatch.setenv('APPLE_CLIENT_ID', 'tw.qingshan.teamparent')
    monkeypatch.setenv('IDENTITY_ENCRYPTION_KEY', Fernet.generate_key().decode())
    monkeypatch.setattr(mobile_api, 'apple_client_secret', lambda: 'test-secret')
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    monkeypatch.setattr(mobile_api, 'apple_keys', SimpleNamespace(get_signing_key_from_jwt=lambda _: SimpleNamespace(key=key.public_key())))
    challenge = client.post('/api/auth/apple/challenge').json()
    claims = {'iss': 'https://appleid.apple.com', 'aud': 'tw.qingshan.teamparent',
              'sub': 'apple-user', 'iat': int(time.time()), 'exp': int(time.time())+300, 'nonce': challenge['nonce']}
    token = jwt.encode(claims, key, algorithm='RS256')
    monkeypatch.setattr(mobile_api, 'exchange_apple_code', lambda _: {'id_token': token, 'refresh_token': 'provider-token'})
    payload = {'challenge': challenge['challenge'], 'identity_token': token, 'authorization_code': 'code'}
    login = client.post('/api/auth/apple', json=payload)
    assert login.status_code == 200
    assert client.post('/api/auth/apple', json=payload).status_code == 401
    with backend.db() as conn, conn.cursor() as cur:
        cur.execute("SELECT encrypted_refresh_token FROM auth_identities WHERE provider='apple'")
        assert cur.fetchone()['encrypted_refresh_token'] != 'provider-token'
    assert client.post('/api/me/deletion-request', json={'confirmation': '刪除帳號'}, headers=headers(login.json()['access_token'])).status_code == 200
    monkeypatch.setattr(apple_identity, 'apple_client_secret', lambda: 'test-secret')
    calls = []
    def fake_revoke(url, data, timeout):
        calls.append(url); assert data['token'] == 'provider-token'
        return SimpleNamespace(status_code=200)
    monkeypatch.setattr(apple_identity.httpx, 'post', fake_revoke)
    assert apple_identity.drain_revocations(backend.db) == 1
    assert calls == ['https://appleid.apple.com/auth/revoke']
    with backend.db() as conn, conn.cursor() as cur:
        cur.execute("SELECT status,encrypted_token FROM provider_revocations ORDER BY id DESC LIMIT 1")
        row = cur.fetchone()
        assert row['status'] == 'completed' and row['encrypted_token'] == ''


def test_simultaneous_refresh_replay_revokes_rotated_session(seeded):
    client, _ = seeded
    session = issue_session(backend.db, 1)
    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(lambda _: client.post('/api/auth/refresh',
            json={'refresh_token': session['refresh_token']}), range(2)))
    assert sorted(r.status_code for r in responses) == [200, 401]
    rotated = next(r.json() for r in responses if r.status_code == 200)
    assert client.get('/api/me', headers=headers(rotated['access_token'])).status_code == 401


def test_admin_input_errors_are_422_without_database_cast_failure(seeded):
    client, tokens = seeded
    event = {'title': 'Practice', 'location': 'park', 'event_date': '2026-02-30', 'player_ids': [1]}
    assert client.post('/api/admin/events', json=event, headers=headers(tokens[0])).status_code == 422
    event.update(event_date='2026-10-02', player_ids=[])
    assert client.post('/api/admin/events', json=event, headers=headers(tokens[0])).status_code == 422
    assert client.post('/api/admin/payments', headers=headers(tokens[0]), json={
        'player_id': 1, 'title': 'fee', 'amount': 500, 'due_date': '2026-02-30'}).status_code == 422
